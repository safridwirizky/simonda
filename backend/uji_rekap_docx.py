"""Menguji pembuat rekap Word (inovasi/rekap_docx.py): perapian teks narasi
hasil salin-tempel dan struktur berkas .docx yang dihasilkan. Jalankan tanpa
Django:

    python uji_rekap_docx.py
"""
import io
import zipfile
from datetime import date

from docx import Document
from docx.oxml.ns import qn

from inovasi import rekap_docx as rd

lolos = gagal = 0


def cek(nama, syarat, tambahan=""):
    global lolos, gagal
    if syarat:
        lolos += 1
        print(f"  OK   {nama}")
    else:
        gagal += 1
        print(f"  GAGAL {nama} {tambahan}")


def ringkas(teks, bagian=""):
    return [(b.jenis, b.teks) for b in rd.susun_blok(teks, bagian)]


print("\n== Perapian teks ==")
terpotong = (
    "Sistem ini dirancang sebagai solusi terhadap pengelolaan arsip yang masih dilakukan\n"
    "secara konvensional, seperti penyimpanan dokumen dalam bentuk fisik, kesulitan\n"
    "menemukan kembali dokumen, serta risiko kerusakan atau kehilangan arsip.\n"
    "\n"
    "Paragraf kedua berdiri sendiri."
)
hasil = ringkas(terpotong)
cek("baris terpotong paksa disambung jadi satu paragraf", len(hasil) == 2 and hasil[0][0] == "p", hasil)
cek("tidak ada pemutus baris tersisa", all("\n" not in t for _, t in hasil))

hasil = ringkas("Tujuan:\n1. Meningkatkan efektivitas\n2.Mempermudah pencarian\n3) Mengurangi risiko")
cek("nomor manual diubah menjadi butir 'num'",
    [j for j, _ in hasil] == ["p", "num", "num", "num"], hasil)
cek("label nomor dibuang dari teks butir", hasil[1][1] == "Meningkatkan efektivitas", hasil[1])

hasil = rd.susun_blok("1. Butir pertama yang panjang dan terpotong\npaksa di sini.\n2. Butir kedua.")
cek("butir bernomor yang terpotong disambung",
    hasil[0].teks == "Butir pertama yang panjang dan terpotong paksa di sini." and hasil[1].nomor == 2,
    [(b.nomor, b.teks) for b in hasil])

hasil = ringkas("•\tDatabase angkatan kerja\n•\tProfil skill\n- Dashboard")
cek("bullet (•, tab, tanda hubung) menjadi 'bul'", [j for j, _ in hasil] == ["bul"] * 3, hasil)

hasil = ringkas("Gambaran Umum Sistem\n\nSistem ini dibangun untuk mendukung pengelolaan arsip digital.")
cek("baris pendek tanpa titik yang disusul isi jadi sub-judul", hasil[0] == ("sub", "Gambaran Umum Sistem"), hasil)

hasil = ringkas("Efesiensi Waktu")
cek("isi satu baris pendek tetap paragraf, bukan sub-judul", hasil == [("p", "Efesiensi Waktu")], hasil)

hasil = ringkas("TUJUAN INOVASI\nMeningkatkan layanan.", "Tujuan")
cek("pengulangan judul bagian di baris pertama dibuang", hasil == [("p", "Meningkatkan layanan.")], hasil)

hasil = ringkas("Hasil Inovasi:\nTersedia 1 aplikasi.", "Dampak/Hasil")
cek("'Hasil Inovasi:' dibuang dari bagian Dampak/Hasil", hasil == [("p", "Tersedia 1 aplikasi.")], hasil)

hasil = ringkas("Penerapan Si-\nVERA memangkas waktu verifikasi.")
cek("kata bertanda hubung yang terpotong disambung tanpa spasi",
    hasil == [("p", "Penerapan Si-VERA memangkas waktu verifikasi.")], hasil)

hasil = ringkas("a. Jangka Pendek :\n1. Satu\n2. Dua\nb. Jangka Menengah :\n1. Tiga")
cek("label huruf pendek jadi sub-judul dan daftar bernomor dimulai ulang",
    [j for j, _ in hasil] == ["sub", "num", "num", "sub", "num"], hasil)

hasil = ringkas("Dokumen diterima\n↓\nDigitalisasi dokumen\n↓\nPenyimpanan")
cek("alur berpanah menjadi blok 'alur', bukan sub-judul",
    all(j == "alur" for j, _ in hasil) and len(hasil) == 5, hasil)

hasil = ringkas("a. Meningkatkan efektivitas dan efisiensi proses monitoring tindak lanjut\nusulan masyarakat.\nb. Kedua.")
cek("butir berhuruf panjang jadi daftar huruf (bukan sub-judul)",
    [j for j, _ in hasil] == ["abj", "abj"], hasil)

hasil = ringkas("Paragraf satu panjang sekali yang berdiri sendiri dan diakhiri titik di ujung baris.\n"
                "Paragraf dua juga panjang dan berdiri sendiri pada baris yang berbeda dari sumber.")
cek("dua baris utuh bertitik tidak digabung", len(hasil) == 2, hasil)

cek("teks kosong menghasilkan nol blok", rd.susun_blok("  \n\n ") == [])
cek("None aman", rd.susun_blok(None) == [])

print("\n== Struktur berkas .docx ==")
contoh = [
    rd.ItemRekap(
        nama="SIRINDU – Sistem Informasi Reminder", opd="DISPUSIP", status="diajukan",
        status_label="Menunggu verifikasi",
        rancang_bangun="Sistem ini dibangun untuk mengelola data.\n\nFitur utama:\n• Reminder\n• Dashboard",
        tujuan="1. Meningkatkan layanan\n2. Mempercepat proses",
        manfaat="Mempermudah akses dokumen",
        hasil="",
    ),
    rd.ItemRekap(
        nama="SIPEKA", opd="PKM_KORBAFO", status="draft", status_label="Draft",
        rancang_bangun="Rancang bangun SIPEKA.",
        tujuan="1. Tujuan satu\n2. Tujuan dua",
        manfaat="Manfaat.", hasil="Hasil.",
    ),
]
data = rd.buat_rekap_docx(contoh, tanggal=date(2026, 10, 22), tahapan="Uji Coba")
cek("hasil berupa zip .docx yang utuh", zipfile.ZipFile(io.BytesIO(data)).testzip() is None)
dok = Document(io.BytesIO(data))
sec = dok.sections[0]

cek("ukuran halaman A4", (sec.page_width.twips, sec.page_height.twips) == (11906, 16838),
    (sec.page_width.twips, sec.page_height.twips))
cek("margin 1440/1440/1300/1440",
    (sec.top_margin.twips, sec.right_margin.twips, sec.bottom_margin.twips, sec.left_margin.twips)
    == (1440, 1440, 1300, 1440))
teks_dok = [p.text for p in dok.paragraphs]
cek("judul dokumen kapital", teks_dok[0] == "REKAPITULASI RANCANG BANGUN, TUJUAN, MANFAAT, DAN HASIL INOVASI DAERAH")
cek("sub-judul memuat wilayah dan tahapan", teks_dok[1] == "Kabupaten Rote Ndao – Inovasi Uji Coba", teks_dok[1])
cek("tanggal berbahasa Indonesia", teks_dok[2] == "Diunduh 22 Oktober 2026 · 2 inovasi", teks_dok[2])
cek("header memuat nama wilayah", "Kabupaten Rote Ndao" in sec.header.paragraphs[0].text)
pie_xml = sec.footer._element.xml
cek("footer memuat field PAGE dan NUMPAGES", "PAGE" in pie_xml and "NUMPAGES" in pie_xml)

tabel = dok.tables[0]
cek("tabel daftar: 1 baris judul + 2 baris data, 4 kolom", len(tabel.rows) == 3 and len(tabel.columns) == 4)
cek("baris judul tabel berulang di tiap halaman", tabel.rows[0]._tr.xpath("./w:trPr/w:tblHeader") != [])
cek("baris data tidak terbelah antarhalaman", tabel.rows[1]._tr.xpath("./w:trPr/w:cantSplit") != [])
cek("lebar kolom 600/3800/2700/1926",
    [int(g.get(qn("w:w"))) for g in tabel._tbl.tblGrid.findall(qn("w:gridCol"))] == [600, 3800, 2700, 1926])
cek("isi sel tabel benar",
    [c.text for c in tabel.rows[1].cells] == ["1", "SIRINDU – Sistem Informasi Reminder", "DISPUSIP", "Menunggu verifikasi"])
cek("baris genap diberi warna selang-seling", tabel.rows[2]._tr.xpath(".//w:shd[@w:fill='F5F7FA']") != [])

h1 = [p for p in dok.paragraphs if p.style.name == "Heading 1"]
h2 = [p for p in dok.paragraphs if p.style.name == "Heading 2"]
cek("satu Heading 1 per inovasi, bernomor", [p.text for p in h1] == ["1. SIRINDU – Sistem Informasi Reminder", "2. SIPEKA"],
    [p.text for p in h1])
cek("Heading 1 pertama mulai di halaman baru, berikutnya tidak",
    h1[0].paragraph_format.page_break_before is True and not h1[1].paragraph_format.page_break_before)
cek("tiap inovasi punya 4 Heading 2 bagian (+1 'Daftar Inovasi')", len(h2) == 1 + 4 * len(contoh), len(h2))
cek("urutan Heading 2 per inovasi", [p.text for p in h2[1:5]] == ["Rancang Bangun", "Tujuan", "Manfaat", "Dampak/Hasil"])

gaya_xml = dok.styles.element.xml
cek("font Heading memakai Arial, bukan tema Calibri bawaan", "asciiTheme" not in gaya_xml and "majorHAnsi" not in gaya_xml)
cek("tanpa jarak/spasi bawaan template di docDefaults", "w:after=\"200\"" not in gaya_xml.split("</w:docDefaults>")[0])

bagian_hasil = []
ambil = False
for p in dok.paragraphs:
    if p.style.name == "Heading 2" and p.text == "Dampak/Hasil":
        ambil = True
        continue
    if ambil:
        bagian_hasil.append(p.text)
        break
cek("bagian kosong diisi tanda '-'", bagian_hasil == ["-"], bagian_hasil)

daftar = [p for p in dok.paragraphs if p.style.name == "List Paragraph"]
cek("butir daftar memakai numPr asli (bukan ketikan manual)", daftar and all(p._p.xpath("./w:pPr/w:numPr") for p in daftar))
cek("tidak ada awalan '1.' / bullet manual di teks butir",
    not any(p.text[:2] in ("1.", "2.", "• ") for p in daftar), [p.text for p in daftar])

numbering = dok.part.numbering_part.element
anak = [c.tag for c in numbering]
terakhir_abstrak = max(i for i, t in enumerate(anak) if t == qn("w:abstractNum"))
pertama_num = min(i for i, t in enumerate(anak) if t == qn("w:num"))
cek("semua abstractNum mendahului num (syarat skema OOXML)", terakhir_abstrak < pertama_num)

id_daftar_nomor = [p._p.xpath("./w:pPr/w:numPr/w:numId/@w:val")[0] for p in daftar if p.text.startswith(("Meningkatkan", "Mempercepat", "Tujuan satu", "Tujuan dua"))]
cek("tiap daftar bernomor memakai instance sendiri agar mulai dari 1",
    len(set(id_daftar_nomor)) == 2, id_daftar_nomor)
id_valid = {n.get(qn("w:numId")) for n in numbering.findall(qn("w:num"))}
cek("semua numId yang dipakai terdaftar di numbering.xml",
    all(p._p.xpath("./w:pPr/w:numPr/w:numId/@w:val")[0] in id_valid for p in daftar))

print(f"\n{lolos} lolos, {gagal} gagal")
raise SystemExit(1 if gagal else 0)
