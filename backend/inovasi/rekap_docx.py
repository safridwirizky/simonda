"""Pembuat dokumen Word (.docx) untuk rekap Rancang Bangun, Tujuan, Manfaat,
dan Dampak/Hasil inovasi daerah.

Modul ini sengaja tidak bergantung pada Django/ORM supaya mudah diuji
(`python uji_rekap_docx.py`). Pemanggilnya (api.py) cukup menyerahkan daftar
`ItemRekap`.

Tata letak mengikuti dokumen rekap yang sudah dirapikan secara manual:
  * A4, margin 2,54 cm, font Arial 11 pt, bahasa Indonesia
  * judul tengah + garis bawah, lalu tabel "Daftar Inovasi" (header berulang,
    baris selang-seling)
  * tiap inovasi: Heading 1 bernomor, baris "OPD | Status" berwarna, lalu
    Heading 2 untuk Rancang Bangun / Tujuan / Manfaat / Dampak/Hasil
  * isi rata kiri-kanan; daftar bernomor dan bullet berupa daftar Word asli
    (bukan ketikan "1." / "•" manual)
  * header dan footer "Halaman X dari Y"

Teks narasi di basis data umumnya hasil salin-tempel (baris terpotong paksa
tiap ~80 karakter, nomor/bullet diketik manual, judul bagian ikut tersalin).
`susun_blok()` merapikannya menjadi paragraf, sub-judul, bullet, dan daftar
bernomor.
"""
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from typing import List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, RGBColor, Twips

# --- palet & ukuran (diambil dari dokumen acuan) ---------------------------
BIRU_TUA = "1F3A5F"
BIRU = "2E6DA4"
ABU = "595959"
GARIS = "BFC7D1"
ZEBRA = "F5F7FA"
FONT = "Arial"

LEBAR_ISI = 9026  # twips: A4 (11906) dikurangi margin kiri+kanan (2 x 1440)
LEBAR_KOLOM = (600, 3800, 2700, 1926)  # No | Nama Inovasi | OPD | Status

WARNA_STATUS = {
    "draft": "7F7F7F",
    "diajukan": "B26A00",
    "revisi": "B42318",
    "terverifikasi": "1E7B34",
}

BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
         "Agustus", "September", "Oktober", "November", "Desember"]

# Judul bagian -> kata yang kalau muncul sebagai baris pertama isi dianggap
# pengulangan judul dan dibuang ("TUJUAN INOVASI", "Manfaat", dst).
_ALIAS_JUDUL = {
    "Rancang Bangun": ["rancang bangun"],
    "Tujuan": ["tujuan"],
    "Manfaat": ["manfaat"],
    "Dampak/Hasil": ["dampak\\s*/\\s*hasil", "dampak", "hasil"],
}


@dataclass
class ItemRekap:
    nama: str
    opd: str
    status: str  # kode status (draft/diajukan/revisi/terverifikasi)
    status_label: str
    rancang_bangun: str = ""
    tujuan: str = ""
    manfaat: str = ""
    hasil: str = ""


def tanggal_indonesia(d: date) -> str:
    return f"{d.day} {BULAN[d.month - 1]} {d.year}"


# ===========================================================================
# 1. Perapian teks narasi
# ===========================================================================
_AKHIR_KALIMAT = re.compile(r"""[.!?:;]["”’')\]]*$""")
_BULLET = re.compile(r"^(?:[•●○▪■◦·]\s*|[-–—*]\s+)(\S.*)$")
_NOMOR = re.compile(r"^(\d{1,2})\s*[.)]\s*([^\d\s].*)$")
_PANAH = re.compile(r"^[↓↑→←⇓⇒⟶▼▲]+$")
_LABEL_HURUF = re.compile(r"^([A-Za-z])[.)]\s+(\S.*)$")  # "a. Jangka Pendek", "B. Bagi masyarakat."


def _bersihkan_baris(teks: str) -> List[str]:
    teks = (teks or "").replace("\r\n", "\n").replace("\r", "\n")
    teks = teks.replace(" ", "\n").replace("\x0b", "\n").replace("\x0c", "\n")
    teks = teks.replace(" ", " ").replace("­", "").replace("\t", " ")
    teks = re.sub("[​-‍﻿]", "", teks)
    # buang karakter kontrol yang tidak boleh ada di XML Word
    teks = re.sub("[\x00-\x08\x0e-\x1f]", "", teks)
    return [re.sub(r" {2,}", " ", baris).strip() for baris in teks.split("\n")]


def _buang_judul_ganda(baris: List[str], judul_bagian: str) -> List[str]:
    """Hapus baris pertama bila hanya mengulang nama bagian."""
    alternatif = "|".join(_ALIAS_JUDUL.get(judul_bagian, []))
    if not alternatif:
        return baris
    pola = re.compile(rf"(?:{alternatif})(?:\s+inovasi(?:\s+daerah)?)?\s*[:\-–]?", re.I)
    for i, b in enumerate(baris):
        if not b:
            continue
        if pola.fullmatch(b):
            return baris[i + 1:]
        break
    return baris


class _Blok:
    """Satu unit isi: jenis 'p' (paragraf), 'sub' (sub-judul), 'bul' (bullet),
    'num' (butir 1. 2. 3.), 'abj'/'ABJ' (butir a. b. c. / A. B. C.), atau
    'alur' (langkah alur/panah, rata tengah)."""

    __slots__ = ("jenis", "teks", "nomor", "panjang_terakhir")

    def __init__(self, jenis, teks, nomor=None):
        self.jenis = jenis
        self.teks = teks
        self.nomor = nomor
        self.panjang_terakhir = len(teks)

    def tambah(self, baris: str):
        # "Si-" + "VERA" -> "Si-VERA": kata bertanda hubung yang terpotong di
        # ujung baris disambung tanpa spasi.
        self.teks += baris if self.teks.endswith("-") else " " + baris
        self.panjang_terakhir = len(baris)


def _huruf_kecil(baris: str) -> bool:
    return baris[:1].islower()


def _lanjutan(blok: _Blok, baris: str, lebar: int) -> bool:
    """Apakah `baris` masih kelanjutan `blok` (baris terpotong paksa)?

    `lebar` = perkiraan lebar pemotongan baris pada sumber (0 bila teks tidak
    terpotong paksa). Baris yang pendek tanpa tanda baca akhir lebih mungkin
    judul/butir berdiri sendiri daripada potongan kalimat."""
    if blok.teks.endswith("-"):
        return True
    selesai = bool(_AKHIR_KALIMAT.search(blok.teks))
    kecil = _huruf_kecil(baris)
    panjang = blok.panjang_terakhir
    if not selesai:
        # tanpa perkiraan lebar (teks pendek), baris >= 60 huruf yang belum
        # ditutup tanda baca dianggap terpotong paksa
        return kecil or panjang >= (0.55 * lebar if lebar else 60)
    if blok.jenis != "p":
        return kecil
    # Kalimat sudah selesai di ujung baris: lanjut hanya bila baris itu hampir
    # sepenuh lebar potong. Baris yang jauh lebih panjang dari lebar potong
    # bukan hasil pemotongan paksa, melainkan satu paragraf utuh.
    return kecil or bool(lebar and 0.8 * lebar <= panjang <= 1.1 * lebar)


def _mirip_sub_judul(baris: str, sisa: List[str]) -> bool:
    """Baris pendek tanpa koma/tanda baca akhir yang disusul isi -> sub-judul
    (mis. "Gambaran Umum E-Arsip BAGORON")."""
    if not [b for b in sisa if b]:  # satu-satunya baris: itu isinya, bukan judul
        return False
    if len(baris) > 60 or len(baris.split()) > 8 or not baris[:1].isalnum() or _huruf_kecil(baris):
        return False
    if baris.endswith("-"):  # kata terpotong ("Si-" / "VERA"), bukan judul
        return False
    if re.search(r"[.,;:]$", baris) or "," in baris or re.fullmatch(r"[\d\s.,%/-]+", baris):
        return False
    langsung = sisa[0] if sisa else ""
    if langsung and _huruf_kecil(langsung):
        return False  # awal kalimat yang terpotong, bukan judul
    return True


def _label_huruf_sub_judul(baris: str, sisa: List[str]) -> bool:
    """"a. Jangka Pendek (6 bulan) :" atau "C. Bagi masyarakat." berdiri
    sebagai sub-judul bila pendek dan tidak disambung baris huruf kecil;
    selain itu ia butir berhuruf biasa ("a. Meningkatkan efektivitas ...")."""
    if len(baris) > 60 or "," in baris or not [b for b in sisa if b]:
        return False
    return not (sisa and sisa[0] and _huruf_kecil(sisa[0]))


def susun_blok(teks: str, judul_bagian: str = "") -> List[_Blok]:
    baris = _buang_judul_ganda(_bersihkan_baris(teks), judul_bagian)

    # Perkirakan lebar pemotongan paksa: persentil ke-80 panjang baris di
    # antara baris panjang yang tidak ditutup tanda baca akhir kalimat
    # (persentil, bukan maksimum, agar satu baris panjang tidak menggeser
    # perkiraan).
    terpotong = sorted(len(b) for b in baris
                       if len(b) >= 40 and not _AKHIR_KALIMAT.search(b))
    lebar = terpotong[int(0.8 * (len(terpotong) - 1))] if len(terpotong) >= 3 else 0

    hasil: List[_Blok] = []
    buf: Optional[_Blok] = None

    def tutup():
        nonlocal buf
        if buf is not None:
            hasil.append(buf)
            buf = None

    for i, b in enumerate(baris):
        sisa = baris[i + 1:]
        if not b:
            tutup()
            continue
        if _PANAH.match(b):
            tutup()
            hasil.append(_Blok("alur", b))
            continue
        m = _NOMOR.match(b)
        if m:
            tutup()
            buf = _Blok("num", m.group(2).strip(), int(m.group(1)))
            continue
        m = _BULLET.match(b)
        if m:
            tutup()
            buf = _Blok("bul", m.group(1).strip())
            continue
        m = _LABEL_HURUF.match(b)
        if m:
            tutup()
            if _label_huruf_sub_judul(b, sisa):
                hasil.append(_Blok("sub", b))
            else:
                huruf = m.group(1)
                buf = _Blok("abj" if huruf.islower() else "ABJ", m.group(2).strip(),
                            ord(huruf.lower()) - 96)
            continue
        if buf is not None and _lanjutan(buf, b, lebar):
            buf.tambah(b)
            continue
        tutup()
        berikutnya_panah = bool(sisa) and bool(_PANAH.match(sisa[0]))
        sebelumnya_alur = bool(hasil) and hasil[-1].jenis == "alur"
        if berikutnya_panah or sebelumnya_alur:
            hasil.append(_Blok("alur", b))
        elif _mirip_sub_judul(b, sisa):
            hasil.append(_Blok("sub", b))
        else:
            buf = _Blok("p", b)
    tutup()
    return hasil


# ===========================================================================
# 2. Gaya dokumen
# ===========================================================================
def _fonts_xml() -> str:
    return (f'<w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}" '
            f'w:eastAsia="{FONT}" w:cs="{FONT}"/>')


def _ganti_isi_gaya(gaya, ppr_xml: str = "", rpr_xml: str = ""):
    """Ganti total pPr/rPr sebuah style dengan XML yang kita tentukan, supaya
    sisa-sisa tema bawaan template python-docx (font Calibri Light, biru
    tema) tidak ikut terbawa."""
    el = gaya.element
    for tag in ("w:pPr", "w:rPr"):
        for lama in el.findall(qn(tag)):
            el.remove(lama)
    if ppr_xml:
        el.append(parse_xml(f"<w:pPr {nsdecls('w')}>{ppr_xml}</w:pPr>"))
    if rpr_xml:
        el.append(parse_xml(f"<w:rPr {nsdecls('w')}>{rpr_xml}</w:rPr>"))


def _siapkan_gaya(dok):
    gaya = dok.styles
    # Template bawaan python-docx memberi jarak 10 pt setelah tiap paragraf dan
    # spasi 1,15 lewat docDefaults; itu yang membuat sel tabel dan header
    # tampak renggang. Dokumen acuan tidak memakainya, jadi dikosongkan dan
    # jarak diatur per paragraf.
    root = gaya.element
    for lama in root.findall(qn("w:docDefaults")):
        root.remove(lama)
    root.insert(0, parse_xml(
        f'<w:docDefaults {nsdecls("w")}><w:rPrDefault><w:rPr>{_fonts_xml()}'
        '<w:sz w:val="22"/><w:szCs w:val="22"/>'
        '<w:lang w:val="id-ID" w:eastAsia="id-ID" w:bidi="ar-SA"/></w:rPr></w:rPrDefault>'
        '<w:pPrDefault/></w:docDefaults>'
    ))
    _ganti_isi_gaya(
        gaya["Normal"],
        rpr_xml=(f'{_fonts_xml()}<w:sz w:val="22"/><w:szCs w:val="22"/>'
                 '<w:lang w:val="id-ID" w:eastAsia="id-ID" w:bidi="ar-SA"/>'),
    )
    _ganti_isi_gaya(
        gaya["Heading 1"],
        ppr_xml=('<w:keepNext/><w:keepLines/>'
                 f'<w:pBdr><w:bottom w:val="single" w:color="{BIRU_TUA}" w:sz="6" w:space="4"/></w:pBdr>'
                 '<w:spacing w:before="440" w:after="80"/><w:outlineLvl w:val="0"/>'),
        rpr_xml=(f'{_fonts_xml()}<w:b/><w:bCs/><w:color w:val="{BIRU_TUA}"/>'
                 '<w:sz w:val="28"/><w:szCs w:val="28"/>'),
    )
    _ganti_isi_gaya(
        gaya["Heading 2"],
        ppr_xml=('<w:keepNext/><w:keepLines/>'
                 '<w:spacing w:before="240" w:after="100"/><w:outlineLvl w:val="1"/>'),
        rpr_xml=(f'{_fonts_xml()}<w:b/><w:bCs/><w:color w:val="{BIRU}"/>'
                 '<w:sz w:val="24"/><w:szCs w:val="24"/>'),
    )
    # Tanpa indentasi/contextualSpacing bawaan: jarak dan indentasi daftar
    # diatur oleh definisi penomoran dan format paragraf kita sendiri.
    _ganti_isi_gaya(gaya["List Paragraph"])
    # Gaya lain bawaan template (Title, Heading 3-9, ...) tidak dipakai, tetapi
    # bila pengguna menerapkannya di Word jangan sampai kembali ke font tema.
    for rf in root.iter(qn("w:rFonts")):
        for atribut in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            rf.attrib.pop(qn(f"w:{atribut}"), None)


def _siapkan_halaman(dok, wilayah: str):
    sec = dok.sections[0]
    sec.page_width, sec.page_height = Twips(11906), Twips(16838)
    sec.top_margin, sec.bottom_margin = Twips(1440), Twips(1300)
    sec.left_margin, sec.right_margin = Twips(1440), Twips(1440)
    sec.header_distance = sec.footer_distance = Twips(708)

    sec.header.is_linked_to_previous = False
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _garis_bawah(hp, GARIS, sz=4, space=4)
    _jalan(hp, f"Rekapitulasi Inovasi Daerah – {wilayah}", ukuran=8, warna=ABU)

    sec.footer.is_linked_to_previous = False
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _jalan(fp, "Halaman ", ukuran=8, warna=ABU)
    _kolom_field(fp, "PAGE")
    _jalan(fp, " dari ", ukuran=8, warna=ABU)
    _kolom_field(fp, "NUMPAGES")


# ===========================================================================
# 3. Penomoran (daftar Word asli)
# ===========================================================================
class _Penomoran:
    """Mendaftarkan definisi bullet, desimal, dan abjad, serta membuat instance
    daftar baru (supaya tiap daftar mulai dari angkanya sendiri)."""

    def __init__(self, dok):
        self.root = dok.part.numbering_part.element
        abstrak = [int(a.get(qn("w:abstractNumId"))) for a in self.root.findall(qn("w:abstractNum"))]
        nums = [int(n.get(qn("w:numId"))) for n in self.root.findall(qn("w:num"))]
        self._abstrak = max(abstrak, default=0) + 1
        self._num = max(nums, default=0) + 1
        self.id_abstrak_bullet = self._tambah_abstrak("bullet", "\u2022", 567, 283)
        self.num_bullet = self._tambah_num(self.id_abstrak_bullet)
        self._abstrak_berurut = {
            "num": self._tambah_abstrak("decimal", "%1.", 567, 340),
            "abj": self._tambah_abstrak("lowerLetter", "%1.", 567, 340),
            "ABJ": self._tambah_abstrak("upperLetter", "%1.", 567, 340),
        }

    def _tambah_abstrak(self, fmt, teks, kiri, gantung) -> int:
        aid = self._abstrak
        self._abstrak += 1
        rpr = (f'<w:rPr><w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}" w:cs="{FONT}" w:hint="default"/></w:rPr>'
               if fmt == "bullet" else "")
        el = parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{aid}">'
            '<w:multiLevelType w:val="hybridMultilevel"/>'
            f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>'
            f'<w:lvlText w:val="{teks}"/><w:lvlJc w:val="left"/>'
            f'<w:pPr><w:ind w:left="{kiri}" w:hanging="{gantung}"/></w:pPr>{rpr}</w:lvl>'
            '</w:abstractNum>'
        )
        # semua abstractNum harus mendahului <w:num> pertama
        pertama_num = self.root.find(qn("w:num"))
        if pertama_num is not None:
            pertama_num.addprevious(el)
        else:
            self.root.append(el)
        return aid

    def _tambah_num(self, id_abstrak: int, mulai: int = 1) -> int:
        nid = self._num
        self._num += 1
        self.root.append(parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{nid}"><w:abstractNumId w:val="{id_abstrak}"/>'
            f'<w:lvlOverride w:ilvl="0"><w:startOverride w:val="{mulai}"/></w:lvlOverride></w:num>'
        ))
        return nid

    def daftar_baru(self, jenis: str, mulai: int = 1) -> int:
        """Instance daftar berurut baru ('num' 1. 2. 3., 'abj' a. b. c.,
        'ABJ' A. B. C.) agar tiap daftar mulai dari angka/hurufnya sendiri."""
        return self._tambah_num(self._abstrak_berurut[jenis], mulai)


def _pasang_num(paragraf, num_id: int):
    ppr = paragraf._p.get_or_add_pPr()
    ppr.append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>'
    ))
    # numPr harus berada tepat setelah pStyle/keepNext/... dan sebelum spacing/jc
    _urutkan_ppr(ppr)


_URUTAN_PPR = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
    "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
    "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
    "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
    "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle",
    "rPr", "sectPr", "pPrChange",
]


def _urutkan_ppr(ppr):
    """Susun ulang anak <w:pPr> sesuai urutan skema OOXML (Word menolak/
    memperbaiki berkas bila urutannya salah)."""
    peringkat = {qn(f"w:{n}"): i for i, n in enumerate(_URUTAN_PPR)}
    anak = list(ppr)
    anak.sort(key=lambda e: peringkat.get(e.tag, len(peringkat)))
    for e in anak:
        ppr.remove(e)
    for e in anak:
        ppr.append(e)


# ===========================================================================
# 4. Pembantu tingkat paragraf / run
# ===========================================================================
def _jalan(paragraf, teks, *, tebal=False, miring=False, ukuran=11, warna=None):
    r = paragraf.add_run(teks)
    r.bold = tebal or None
    r.italic = miring or None
    r.font.size = Pt(ukuran)
    if warna:
        r.font.color.rgb = RGBColor.from_string(warna)
    return r


def _garis_bawah(paragraf, warna, sz, space):
    ppr = paragraf._p.get_or_add_pPr()
    ppr.append(parse_xml(
        f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:color="{warna}" '
        f'w:sz="{sz}" w:space="{space}"/></w:pBdr>'
    ))
    _urutkan_ppr(ppr)


def _kolom_field(paragraf, instruksi: str):
    """Field PAGE/NUMPAGES dengan teks pengganti "1" supaya tetap tampil di
    penampil yang tidak memperbarui field."""
    def _run():
        r = paragraf.add_run()
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(ABU)
        return r

    def _fld(r, tipe):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), tipe)
        r._r.append(el)

    r1 = _run(); _fld(r1, "begin")
    r2 = _run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instruksi} "
    r2._r.append(it)
    r3 = _run(); _fld(r3, "separate")
    r4 = _run(); r4.text = "1"
    r5 = _run(); _fld(r5, "end")


def _atur(paragraf, *, sebelum=None, sesudah=None, baris=None, rata=None, keep_next=False):
    pf = paragraf.paragraph_format
    if sebelum is not None:
        pf.space_before = Twips(sebelum)
    if sesudah is not None:
        pf.space_after = Twips(sesudah)
    if baris is not None:
        # satuan 1/240 baris (300 = 1,25 baris); ditulis langsung agar tidak
        # dibulatkan oleh konversi pecahan python-docx
        sp = paragraf._p.get_or_add_pPr().get_or_add_spacing()
        sp.set(qn("w:line"), str(baris))
        sp.set(qn("w:lineRule"), "auto")
    if rata is not None:
        paragraf.alignment = rata
    if keep_next:
        pf.keep_with_next = True


def _paragraf_isi(dok, teks, *, num_id=None):
    p = dok.add_paragraph(style="List Paragraph" if num_id else None)
    if num_id:
        _atur(p, sesudah=80, baris=290, rata=WD_ALIGN_PARAGRAPH.JUSTIFY)
        _pasang_num(p, num_id)
    else:
        _atur(p, sesudah=120, baris=300, rata=WD_ALIGN_PARAGRAPH.JUSTIFY)
    _jalan(p, teks, ukuran=11)
    return p


def _tulis_blok(dok, blok: List[_Blok], nomor: _Penomoran):
    num_aktif: Optional[int] = None  # instance daftar desimal yang sedang berjalan
    for idx, b in enumerate(blok):
        if b.jenis not in ("num", "abj", "ABJ"):
            num_aktif = None
        if b.jenis == "p":
            _paragraf_isi(dok, b.teks)
        elif b.jenis == "sub":
            p = dok.add_paragraph()
            _atur(p, sebelum=160, sesudah=80, keep_next=True)
            _jalan(p, b.teks, tebal=True, ukuran=11, warna=BIRU_TUA)
        elif b.jenis == "bul":
            _paragraf_isi(dok, b.teks, num_id=nomor.num_bullet)
        elif b.jenis in ("num", "abj", "ABJ"):
            sebelumnya = blok[idx - 1] if idx else None
            if num_aktif is None or sebelumnya is None or sebelumnya.jenis != b.jenis:
                # daftar baru: pertahankan angka/huruf awalnya (mis. daftar yang
                # terputus paragraf dan dilanjutkan dari "2.")
                num_aktif = nomor.daftar_baru(b.jenis, b.nomor or 1)
            _paragraf_isi(dok, b.teks, num_id=num_aktif)
        elif b.jenis == "alur":
            p = dok.add_paragraph()
            _atur(p, sesudah=60, rata=WD_ALIGN_PARAGRAPH.CENTER)
            _jalan(p, b.teks, ukuran=11)


# ===========================================================================
# 5. Tabel "Daftar Inovasi"
# ===========================================================================
def _sel(sel, lebar, teks, *, rata, tebal=False, warna=None, isi=None):
    tcpr = sel._tc.get_or_add_tcPr()
    for anak in list(tcpr):
        tcpr.remove(anak)
    garis = "".join(
        f'<w:{sisi} w:val="single" w:color="{GARIS}" w:sz="4"/>'
        for sisi in ("top", "left", "bottom", "right")
    )
    tcpr.append(parse_xml(f'<w:tcW {nsdecls("w")} w:type="dxa" w:w="{lebar}"/>'))
    tcpr.append(parse_xml(f'<w:tcBorders {nsdecls("w")}>{garis}</w:tcBorders>'))
    if isi:
        tcpr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{isi}"/>'))
    tcpr.append(parse_xml(
        f'<w:tcMar {nsdecls("w")}><w:top w:w="70" w:type="dxa"/><w:left w:w="100" w:type="dxa"/>'
        '<w:bottom w:w="70" w:type="dxa"/><w:right w:w="100" w:type="dxa"/></w:tcMar>'
    ))
    tcpr.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))
    p = sel.paragraphs[0]
    p.alignment = rata
    _jalan(p, teks, tebal=tebal, ukuran=9.5, warna=warna)


def _tabel_daftar(dok, items: List[ItemRekap]):
    tabel = dok.add_table(rows=1 + len(items), cols=4)
    tbl = tabel._tbl
    tbl_pr = tbl.tblPr
    for anak in list(tbl_pr):
        tbl_pr.remove(anak)
    garis = "".join(
        f'<w:{sisi} w:val="single" w:color="{GARIS}" w:sz="4"/>'
        for sisi in ("top", "left", "bottom", "right", "insideH", "insideV")
    )
    tbl_pr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:type="dxa" w:w="{LEBAR_ISI}"/>'))
    tbl_pr.append(parse_xml(f'<w:tblBorders {nsdecls("w")}>{garis}</w:tblBorders>'))
    tbl_pr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
    for kolom, lebar in zip(tbl.tblGrid.findall(qn("w:gridCol")), LEBAR_KOLOM):
        kolom.set(qn("w:w"), str(lebar))

    kepala = ("No", "Nama Inovasi", "OPD", "Status")
    rata = (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT)
    baris0 = tabel.rows[0]
    baris0._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for sel, lebar, teks, r in zip(baris0.cells, LEBAR_KOLOM, kepala, rata):
        _sel(sel, lebar, teks, rata=r, tebal=True, warna="FFFFFF", isi=BIRU_TUA)

    for n, (baris, item) in enumerate(zip(tabel.rows[1:], items), start=1):
        baris._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        isi = ZEBRA if n % 2 == 0 else None
        nilai = (str(n), item.nama, item.opd, item.status_label)
        for sel, lebar, teks, r in zip(baris.cells, LEBAR_KOLOM, nilai, rata):
            _sel(sel, lebar, teks, rata=r, isi=isi)


# ===========================================================================
# 6. Perakitan dokumen
# ===========================================================================
def buat_rekap_docx(
    items: List[ItemRekap],
    *,
    tanggal: Optional[date] = None,
    wilayah: str = "Kabupaten Rote Ndao",
    tahapan: Optional[str] = None,
) -> bytes:
    """Bangun berkas .docx rekap dan kembalikan isinya sebagai bytes.

    `tahapan`: bila seluruh inovasi satu tahapan (mis. "Uji Coba"), dicantumkan
    pada sub-judul."""
    tanggal = tanggal or date.today()
    dok = Document()
    _siapkan_gaya(dok)
    _siapkan_halaman(dok, wilayah)
    nomor = _Penomoran(dok)

    inti = dok.core_properties
    inti.title = "Rekapitulasi Rancang Bangun, Tujuan, Manfaat, dan Hasil Inovasi Daerah"
    inti.author = "SANDO Rote Ndao"
    inti.last_modified_by = "SANDO Rote Ndao"
    inti.created = inti.modified = datetime.now()
    inti.comments = ""

    # --- blok judul ---
    p = dok.add_paragraph()
    _atur(p, sebelum=200, sesudah=120, rata=WD_ALIGN_PARAGRAPH.CENTER)
    _jalan(p, "REKAPITULASI RANCANG BANGUN, TUJUAN, MANFAAT, DAN HASIL INOVASI DAERAH",
           tebal=True, ukuran=15, warna=BIRU_TUA)

    p = dok.add_paragraph()
    _atur(p, sesudah=60, rata=WD_ALIGN_PARAGRAPH.CENTER)
    subjudul = wilayah + (f" – Inovasi {tahapan}" if tahapan else "")
    _jalan(p, subjudul, ukuran=11, warna=ABU)

    p = dok.add_paragraph()
    _atur(p, sesudah=300, rata=WD_ALIGN_PARAGRAPH.CENTER)
    _garis_bawah(p, BIRU_TUA, sz=8, space=8)
    _jalan(p, f"Diunduh {tanggal_indonesia(tanggal)} · {len(items)} inovasi",
           ukuran=10, warna=ABU)

    # --- tabel ringkasan ---
    dok.add_paragraph("Daftar Inovasi", style="Heading 2")
    _tabel_daftar(dok, items)

    # --- uraian per inovasi ---
    bagian = (
        ("Rancang Bangun", "rancang_bangun"),
        ("Tujuan", "tujuan"),
        ("Manfaat", "manfaat"),
        ("Dampak/Hasil", "hasil"),
    )
    for n, item in enumerate(items, start=1):
        h1 = dok.add_paragraph(f"{n}. {item.nama}", style="Heading 1")
        if n == 1:
            h1.paragraph_format.page_break_before = True

        meta = dok.add_paragraph()
        _atur(meta, sesudah=120, keep_next=True)
        _jalan(meta, item.opd, tebal=True, ukuran=10, warna=ABU)
        _jalan(meta, "  |  ", ukuran=10, warna="BFBFBF")
        _jalan(meta, item.status_label, tebal=True, ukuran=10,
               warna=WARNA_STATUS.get(item.status, ABU))

        for judul, atribut in bagian:
            dok.add_paragraph(judul, style="Heading 2")
            blok = susun_blok(getattr(item, atribut, "") or "", judul)
            if blok:
                _tulis_blok(dok, blok, nomor)
            else:
                _paragraf_isi(dok, "-")

    buffer = io.BytesIO()
    dok.save(buffer)
    return buffer.getvalue()
